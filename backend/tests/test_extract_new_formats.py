"""Tests for the PPTX/image/OCR extensions to app/rag/extract.py. Real
PPTX/PNG/PDF bytes are built in-memory (python-pptx, Pillow, pymupdf -
all three are already hard dependencies of extract.py itself) rather than
checked-in fixture files, so the test is self-contained and the exact
shapes/pages under test are visible right here. OCR itself
(_ocr_image_bytes) is monkeypatched in every test that would otherwise
need a real tesseract binary - this suite runs on a host that doesn't
have tesseract installed (only the Docker image does, see the
Dockerfile), consistent with how embed_text/Ollama is monkeypatched
elsewhere in this suite rather than requiring the real service.
"""

from io import BytesIO

import pytest

import app.rag.extract as extract_mod
from app.rag.extract import ExtractionError, SUPPORTED_EXTENSIONS, extract


def _make_pptx_bytes(slides: list[dict]) -> bytes:
    """slides: [{"title": str|None, "body": str|None, "table": [[str,...]]|None}]"""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    for spec in slides:
        slide = prs.slides.add_slide(blank_layout)
        if spec.get("title"):
            box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
            box.text_frame.text = spec["title"]
        if spec.get("body"):
            box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
            box.text_frame.text = spec["body"]
        if spec.get("table"):
            rows, cols = len(spec["table"]), len(spec["table"][0])
            table_shape = slide.shapes.add_table(rows, cols, Inches(1), Inches(4), Inches(6), Inches(2))
            for r, row in enumerate(spec["table"]):
                for c, cell_text in enumerate(row):
                    table_shape.table.cell(r, c).text = cell_text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _make_png_bytes(size=(80, 40), color=(255, 255, 255)) -> bytes:
    from PIL import Image

    img = Image.new("RGB", size, color=color)
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _make_pdf_bytes(pages: list[str | None]) -> bytes:
    """pages: text to insert per page, or None for a genuinely blank
    (no text layer) page - triggers the OCR fallback path."""
    import fitz

    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    buf = BytesIO(doc.tobytes())
    doc.close()
    return buf.getvalue()


# --- SUPPORTED_EXTENSIONS / dispatch -----------------------------------------


def test_new_extensions_registered():
    assert {"pptx", "jpg", "jpeg", "png"} <= SUPPORTED_EXTENSIONS


# --- PPTX ---------------------------------------------------------------------


def test_pptx_extracts_text_frames_and_tables():
    content = _make_pptx_bytes(
        [
            {"title": "Bilan 2025", "body": "Chiffre d'affaires en hausse.", "table": [["Mois", "CA"], ["Janvier", "420000"]]},
            {"title": "Conclusion", "body": "Merci."},
        ]
    )

    result = extract("presentation.pptx", content)

    assert len(result.pages) == 2
    assert result.pages[0].page_no == 1
    assert "Bilan 2025" in result.pages[0].text
    assert "Chiffre d'affaires en hausse." in result.pages[0].text
    assert "Mois | CA" in result.pages[0].text
    assert "Janvier | 420000" in result.pages[0].text
    assert result.pages[1].page_no == 2
    assert "Conclusion" in result.pages[1].text


def test_pptx_picture_shape_gets_ocr_text(monkeypatch):
    from pptx import Presentation
    from pptx.util import Inches

    monkeypatch.setattr(extract_mod, "_ocr_image_bytes", lambda b: "texte extrait de l'image")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    png_bytes = _make_png_bytes()
    slide.shapes.add_picture(BytesIO(png_bytes), Inches(1), Inches(1), Inches(2), Inches(2))
    buf = BytesIO()
    prs.save(buf)

    result = extract("slide.pptx", buf.getvalue())

    assert "texte extrait de l'image" in result.pages[0].text


def test_pptx_empty_presentation_raises():
    content = _make_pptx_bytes([{}])  # a slide with no title/body/table at all
    with pytest.raises(ExtractionError):
        extract("empty.pptx", content)


# --- images ---------------------------------------------------------------


def test_image_uses_ocr(monkeypatch):
    monkeypatch.setattr(extract_mod, "_ocr_image_bytes", lambda b: "Bonjour depuis l'image")

    result = extract("scan.png", _make_png_bytes())

    assert result.text == "Bonjour depuis l'image"
    assert result.pages[0].page_no is None
    assert result.pages[0].section is None


def test_image_jpeg_extension_also_dispatches_to_ocr(monkeypatch):
    monkeypatch.setattr(extract_mod, "_ocr_image_bytes", lambda b: "jpeg ocr text")

    result = extract("photo.jpeg", _make_png_bytes())  # content format doesn't matter, OCR is mocked

    assert result.text == "jpeg ocr text"


def test_image_no_ocr_text_found_raises(monkeypatch):
    monkeypatch.setattr(extract_mod, "_ocr_image_bytes", lambda b: "")

    with pytest.raises(ExtractionError):
        extract("blank.png", _make_png_bytes())


# --- PDF: real text pages untouched, scanned pages OCR'd --------------------


def test_pdf_real_text_page_does_not_invoke_ocr(monkeypatch):
    def fail_if_called(content, page_index):
        raise AssertionError("OCR must not run on a page that already has a text layer")

    monkeypatch.setattr(extract_mod, "_ocr_pdf_page", fail_if_called)
    content = _make_pdf_bytes(["Ceci est un vrai texte de page PDF."])

    result = extract("doc.pdf", content)

    assert "Ceci est un vrai texte de page PDF." in result.pages[0].text


def test_pdf_scanned_page_falls_back_to_ocr_for_that_page_only(monkeypatch):
    monkeypatch.setattr(extract_mod, "_ocr_pdf_page", lambda content, page_index: f"OCR page {page_index + 1}")
    content = _make_pdf_bytes(["Texte réel page 1", None])  # page 2 has no text layer

    result = extract("mixed.pdf", content)

    assert len(result.pages) == 2
    assert result.pages[0].page_no == 1
    assert "Texte réel page 1" in result.pages[0].text
    assert result.pages[1].page_no == 2
    assert result.pages[1].text == "OCR page 2"


def test_pdf_fully_scanned_with_unreadable_ocr_raises(monkeypatch):
    monkeypatch.setattr(extract_mod, "_ocr_pdf_page", lambda content, page_index: "")
    content = _make_pdf_bytes([None, None])

    with pytest.raises(ExtractionError):
        extract("scanned.pdf", content)


def test_ocr_image_bytes_never_raises_on_garbage_input():
    # Real _ocr_image_bytes (not monkeypatched) - garbage bytes must not
    # blow up extraction, just yield no text (best-effort OCR contract).
    assert extract_mod._ocr_image_bytes(b"not an image") == ""
