"""Names survive the whole indexing pipeline (extraction -> chunking ->
document_chunks/document_pages), for every supported office format, and
the Workspace content search finds them with the same exact matcher.
Real files are generated in memory (python-docx, openpyxl, PyMuPDF,
python-pptx) - no fixture file, no Ollama (embeddings are stubbed)."""

from io import BytesIO
from types import SimpleNamespace

import pytest

from app.agent.local_files import (
    MAX_SEARCH_FILE_BYTES,
    search_local_files,
    select_searchable_files,
)
from app.models.document import CONFIDENTIALITY_RANK, Document, DocumentChunk, DocumentPage
from app.rag.extract import extract
from app.rag.ingest import chunk_text, ingest_document
from tests.conftest import EMBEDDING_DIM

NAMES = ("Tarek Ben Ali", "Firas Trabelsi")
BODY = "Compte rendu de réunion.\nParticipants : Tarek Ben Ali (RH) et Firas Trabelsi (IT).\nFin."


def _txt():
    return "notes.txt", BODY.encode("utf-8")


def _docx():
    from docx import Document as Docx

    doc = Docx()
    doc.add_heading("Compte rendu", level=1)
    for line in BODY.split("\n"):
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    return "compte_rendu.docx", buffer.getvalue()


def _xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Annuaire"
    ws.append(["Nom", "Service"])
    ws.append(["Tarek Ben Ali", "RH"])
    ws.append(["Firas Trabelsi", "IT"])
    buffer = BytesIO()
    wb.save(buffer)
    return "annuaire.xlsx", buffer.getvalue()


def _pdf():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Participants : Tarek Ben Ali (RH) et Firas Trabelsi (IT).")
    return "rapport.pdf", pdf.tobytes()


def _pptx():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Équipe projet"
    slide.placeholders[1].text = "Tarek Ben Ali et Firas Trabelsi"
    buffer = BytesIO()
    presentation.save(buffer)
    return "equipe.pptx", buffer.getvalue()


FORMATS = [_txt, _docx, _xlsx, _pdf, _pptx]


# --- indexation ----------------------------------------------------------------------


@pytest.mark.parametrize("make_file", FORMATS, ids=lambda f: f.__name__)
def test_extraction_and_chunking_keep_the_names(make_file):
    name, content = make_file()
    extraction = extract(name, content)
    for person in NAMES:
        assert person in extraction.text, f"{person} lost by extraction of {name}"
    chunks = chunk_text(extraction.text)
    for person in NAMES:
        assert any(person in c.text for c in chunks), f"{person} lost by chunking of {name}"


@pytest.mark.parametrize("make_file", [_docx, _xlsx, _pdf], ids=lambda f: f.__name__)
def test_ingested_document_stores_the_names_in_chunks_and_pages(make_file, db, monkeypatch):
    monkeypatch.setattr("app.rag.ingest.embed_texts", lambda texts: [[0.1] * EMBEDDING_DIM for _ in texts])
    name, content = make_file()
    extraction = extract(name, content)
    doc = Document(
        title=name, department="GENERAL", confidentiality="INTERNAL",
        confidentiality_rank=CONFIDENTIALITY_RANK["INTERNAL"], content=extraction.text, source_filename=name,
    )
    db.add(doc)
    db.commit()

    ingest_document(db, doc, pages=extraction.pages)

    chunks = [c.chunk_text for c in db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id)]
    pages = [p.text for p in db.query(DocumentPage).filter(DocumentPage.document_id == doc.id)]
    for person in NAMES:
        assert any(person in c for c in chunks), f"{person} not in any chunk of {name}"
        assert any(person in p for p in pages), f"{person} not in any page of {name}"


# --- Workspace content search ------------------------------------------------------------


def _files():
    return [(name, f"dossier/{name}", content) for name, content in (f() for f in FORMATS)]


@pytest.mark.parametrize("query", ["Tarek", "tarek", "TAREK"])
def test_workspace_search_is_case_insensitive(query):
    hits, unreadable = search_local_files(query, _files())
    assert unreadable == []
    assert {h.name for h in hits} == {"notes.txt", "compte_rendu.docx", "annuaire.xlsx", "rapport.pdf", "equipe.pptx"}
    assert all(h.count >= 1 for h in hits)
    assert any("Tarek" in excerpt for h in hits for _, excerpt in h.excerpts)


def test_workspace_search_finds_another_first_name():
    hits, _ = search_local_files("Firas", _files())
    assert len(hits) == 5


def test_workspace_search_is_exact_tarak_is_not_tarek():
    hits, _ = search_local_files("Tarak", _files())
    assert hits == []


def test_unreadable_workspace_file_is_reported_not_dropped():
    hits, unreadable = search_local_files("Tarek", [("casse.docx", "casse.docx", b"pas un docx"), *_files()[:1]])
    assert [h.name for h in hits] == ["notes.txt"]
    assert [u["name"] for u in unreadable] == ["casse.docx"]


def test_searchable_file_selection_respects_types_and_size():
    entry = lambda name, ext, size: SimpleNamespace(name=name, relative_path=name, extension=ext, size=size)  # noqa: E731
    selected, skipped = select_searchable_files(
        [
            entry("cv.docx", "docx", 10_000),
            entry("photo.png", "png", 10_000),  # not searched (OCR would be needed)
            entry("enorme.pdf", "pdf", MAX_SEARCH_FILE_BYTES + 1),
            entry("../secret.txt", "txt", 10),  # unsafe path, never read
        ]
    )
    assert [e.name for e in selected] == ["cv.docx"]
    assert skipped == [{"name": "enorme.pdf", "reason": "fichier trop volumineux"}]
